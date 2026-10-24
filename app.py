from __future__ import annotations

import io
import tempfile
from datetime import datetime
from pathlib import Path

import streamlit as st
import whisper
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

st.set_page_config(page_title="Транскрибация аудио и видео", page_icon="🎙️", layout="centered")
MODELS = {
    "tiny — быстрее, менее точно": "tiny",
    "base — баланс скорости и качества": "base",
    "small — хорошее качество": "small",
    "medium — высокое качество, медленнее": "medium",
    "large — максимальное качество, очень медленно": "large",
}
LANGUAGES = {
    "Автоопределение": None, "Русский": "ru", "Английский": "en",
    "Украинский": "uk", "Немецкий": "de", "Французский": "fr",
    "Испанский": "es", "Китайский": "zh", "Японский": "ja",
}
FILE_TYPES = ["mp3", "wav", "m4a", "flac", "ogg", "mp4", "mpeg", "mpga", "webm", "mov", "mkv", "avi"]

@st.cache_resource(show_spinner="Загружаю модель Whisper…")
def load_model(name: str):
    return whisper.load_model(name)

def timestamp(seconds: float) -> str:
    hours, remainder = divmod(max(0, int(seconds)), 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"

def make_docx(source_name: str, result: dict, model_name: str, language_name: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    title = doc.add_heading("Транскрибация", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = doc.add_paragraph()
    for label, value in [
        ("Исходный файл", source_name), ("Язык", language_name),
        ("Модель", model_name), ("Создано", datetime.now().astimezone().strftime("%d.%m.%Y %H:%M")),
    ]:
        info.add_run(label + ": ").bold = True
        info.add_run(value + "\n")
    doc.add_heading("Полный текст", level=1)
    full_text = result.get("text", "").strip()
    if full_text:
        for paragraph in full_text.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
    else:
        doc.add_paragraph("Речь не распознана.")
    segments = result.get("segments", [])
    if segments:
        doc.add_heading("Текст с временными метками", level=1)
        for item in segments:
            line = doc.add_paragraph()
            line.add_run(f"[{timestamp(item.get('start', 0))}–{timestamp(item.get('end', 0))}] ").bold = True
            line.add_run(item.get("text", "").strip())
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()

st.title("🎙️ Транскрибация аудио и видео")
st.write("Загрузите запись — приложение распознает речь и подготовит документ Word с полным текстом и временными метками.")
st.info("Обработка выполняется локально на вашем компьютере. Записи не отправляются в облачный сервис.")
uploaded = st.file_uploader("Аудио- или видеофайл", type=FILE_TYPES, help="Для видео извлекается звуковая дорожка через FFmpeg.")
left, right = st.columns(2)
with left:
    model_label = st.selectbox("Модель распознавания", list(MODELS.keys()), index=2)
with right:
    language_label = st.selectbox("Язык записи", list(LANGUAGES.keys()), index=0)
with st.expander("Дополнительная настройка"):
    beam_size = st.slider("Точность поиска (beam size)", 1, 10, 5, help="Большее значение может немного повысить точность, но замедляет обработку.")

if uploaded and st.button("Начать транскрибацию", type="primary", use_container_width=True):
    suffix = Path(uploaded.name).suffix.lower() or ".audio"
    try:
        with tempfile.TemporaryDirectory(prefix="transcriber-") as temp_dir:
            source_path = Path(temp_dir) / f"input{suffix}"
            source_path.write_bytes(uploaded.getvalue())
            model_name = MODELS[model_label]
            model = load_model(model_name)
            with st.spinner("Распознаю речь. На длинных файлах и больших моделях обработка может занять время…"):
                result = model.transcribe(
                    str(source_path), language=LANGUAGES[language_label], task="transcribe",
                    fp16=False, beam_size=beam_size, verbose=False,
                )
            detected = result.get("language", "не определён")
            resolved_language = language_label if LANGUAGES[language_label] else detected
            docx_data = make_docx(uploaded.name, result, model_name, resolved_language)
        st.success("Транскрибация завершена.")
        st.download_button(
            "Скачать документ Word (.docx)", data=docx_data,
            file_name=f"{Path(uploaded.name).stem}_transcript.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary", use_container_width=True,
        )
        with st.expander("Предпросмотр полного текста"):
            st.text_area("Текст", result.get("text", ""), height=300, label_visibility="collapsed")
    except FileNotFoundError:
        st.error("Не найден FFmpeg. Установите его и убедитесь, что команда ffmpeg доступна в PATH.")
    except Exception as exc:
        st.error(f"Не удалось обработать файл: {exc}")

st.caption("Точность зависит от записи и модели. Временные метки даются по фрагментам речи; разделение по говорящим не включено.")
